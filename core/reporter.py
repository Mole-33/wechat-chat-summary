from pathlib import Path
from typing import List, Optional
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from rich.console import Console
from rich.table import Table
from rich.panel import Panel
from rich.text import Text

from core.models import GroupDailySummary, ChatMessage
from config import REPORTS_DIR


class ReportGenerator:
    """报表与可视化生成器，支持控制台打印、文字简报与专业 Excel 报表导出"""

    def __init__(self, output_dir: Optional[Path] = None):
        self.output_dir = output_dir or REPORTS_DIR
        self.output_dir.mkdir(parents=True, exist_ok=True)
        self.console = Console(force_terminal=True, legacy_windows=False)

    def generate_text_brief(self, summary: GroupDailySummary, top_n: int = 10) -> str:
        """生成精炼美观的微信群排行榜文字简报（可直接复制发到群内）"""
        lines = [
            f"📊 【{summary.group_name}】{summary.date_str} 聊天日报",
            "═" * 32,
            f"💬 今日发言总量：{summary.total_messages:,} 条",
            f"👥 今日发言人数：{summary.total_members_spoke} 人",
            f"📝 今日总字数：{summary.total_words:,} 字",
        ]

        if summary.peak_hour is not None:
            lines.append(f"⏰ 最热活跃时段：{summary.peak_hour:02d}:00 - {summary.peak_hour+1:02d}:00 ({summary.peak_hour_count}条)")

        lines.append("═" * 32)
        lines.append(f"🏆 今日发言活跃榜 TOP {min(top_n, len(summary.member_stats))}：")

        medal_map = {1: "🥇", 2: "🥈", 3: "🥉"}
        for stat in summary.member_stats[:top_n]:
            rank_icon = medal_map.get(stat.rank, f" {stat.rank}.")
            pct = f"{stat.ratio * 100:.1f}%"
            lines.append(f"{rank_icon} {stat.nickname}：{stat.message_count} 条 ({pct})")

        if len(summary.member_stats) > top_n:
            lines.append(f"... 还有 {len(summary.member_stats) - top_n} 位群友参与了互动")

        lines.append("═" * 32)
        lines.append("祝大家交流愉快，期待明天的精彩话题！✨")
        return "\n".join(lines)

    def print_console_table(self, summary: GroupDailySummary, top_n: int = 15):
        """在控制台输出高颜值、彩色格式的每日统计面板"""
        title = f"[bold cyan]📊 微信群每日统计分析 - {summary.group_name} ({summary.date_str})[/bold cyan]"

        # 核心指标面板
        peak_info = f"{summary.peak_hour:02d}:00 - {summary.peak_hour+1:02d}:00 ({summary.peak_hour_count}条)" if summary.peak_hour is not None else "暂无"
        kpi_text = (
            f"[bold]总发言条数:[/bold] [yellow]{summary.total_messages:,}[/yellow] 条   |   "
            f"[bold]发言群员:[/bold] [green]{summary.total_members_spoke}[/green] 人   |   "
            f"[bold]总字数:[/bold] [magenta]{summary.total_words:,}[/magenta] 字   |   "
            f"[bold]最热时段:[/bold] [cyan]{peak_info}[/cyan]"
        )
        self.console.print(Panel(kpi_text, title=title, border_style="cyan"))

        # 群员排行榜表格
        table = Table(title=f"🏆 今日群员发言榜 (TOP {min(top_n, len(summary.member_stats))})", border_style="blue")
        table.add_column("排名", justify="center", style="bold")
        table.add_column("群员昵称", style="bold green", min_width=16)
        table.add_column("发言条数", justify="right", style="yellow")
        table.add_column("占比", justify="right", style="cyan")
        table.add_column("总字数", justify="right", style="magenta")
        table.add_column("平均字数/条", justify="right")
        table.add_column("首条时间", justify="center", style="dim")
        table.add_column("末条时间", justify="center", style="dim")

        medal_map = {1: "🥇 1", 2: "🥈 2", 3: "🥉 3"}
        for stat in summary.member_stats[:top_n]:
            rank_str = medal_map.get(stat.rank, str(stat.rank))
            pct = f"{stat.ratio * 100:.1f}%"
            table.add_row(
                rank_str,
                stat.nickname,
                f"{stat.message_count:,}",
                pct,
                f"{stat.total_words:,}",
                f"{stat.avg_words_per_msg}",
                stat.first_msg_time or "-",
                stat.last_msg_time or "-"
            )

        self.console.print(table)

    def export_to_excel(
        self,
        summary: GroupDailySummary,
        messages: Optional[List[ChatMessage]] = None,
        output_path: Optional[Path] = None
    ) -> Path:
        """导出专业格式的 Excel 工作簿，含排行榜、24小时时段分布及全天明细"""
        if output_path is None:
            # 安全文件名过滤
            safe_group = "".join(c for c in summary.group_name if c.isalnum() or c in ('_', '-'))
            filename = f"微信群统计_{safe_group}_{summary.date_str}.xlsx"
            output_path = self.output_dir / filename

        wb = openpyxl.Workbook()

        # 样式定义
        header_font = Font(name="Microsoft YaHei", size=11, bold=True, color="FFFFFF")
        header_fill = PatternFill(start_color="2B579A", end_color="2B579A", fill_type="solid")
        title_font = Font(name="Microsoft YaHei", size=14, bold=True, color="2B579A")
        bold_font = Font(name="Microsoft YaHei", size=11, bold=True)
        regular_font = Font(name="Microsoft YaHei", size=10)
        thin_border = Border(
            left=Side(style="thin", color="D3D3D3"),
            right=Side(style="thin", color="D3D3D3"),
            top=Side(style="thin", color="D3D3D3"),
            bottom=Side(style="thin", color="D3D3D3")
        )
        center_align = Alignment(horizontal="center", vertical="center")
        left_align = Alignment(horizontal="left", vertical="center")
        right_align = Alignment(horizontal="right", vertical="center")

        # -----------------------------------------------------------------
        # Sheet 1: 群员发言排行榜与每日概览
        # -----------------------------------------------------------------
        ws1 = wb.active
        ws1.title = "今日排行与概览"
        ws1.views.sheetView[0].showGridLines = True

        # 标题栏
        ws1.merge_cells("A1:G1")
        ws1["A1"] = f"【{summary.group_name}】每日聊天统计报表 - {summary.date_str}"
        ws1["A1"].font = title_font
        ws1["A1"].alignment = center_align

        # 核心 KPI 汇总卡片
        kpis = [
            ("总发言条数", f"{summary.total_messages:,} 条"),
            ("活跃群员数", f"{summary.total_members_spoke} 人"),
            ("全天总字数", f"{summary.total_words:,} 字"),
            ("最热时段", f"{summary.peak_hour:02d}:00 - {summary.peak_hour+1:02d}:00 ({summary.peak_hour_count}条)" if summary.peak_hour is not None else "-")
        ]
        ws1.append([])  # 空行
        for idx, (label, val) in enumerate(kpis, 1):
            col_letter = get_column_letter(idx * 2 - 1)
            next_col = get_column_letter(idx * 2)
            ws1[f"{col_letter}3"] = label
            ws1[f"{col_letter}3"].font = regular_font
            ws1[f"{col_letter}3"].alignment = center_align
            ws1[f"{col_letter}4"] = val
            ws1[f"{col_letter}4"].font = bold_font
            ws1[f"{col_letter}4"].alignment = center_align

        ws1.append([])  # 空行 (row 5)

        # 排行榜表头
        headers1 = ["名次", "群员昵称", "发言条数", "条数占比", "总发言字数", "首条发言时间", "最后发言时间"]
        ws1.append(headers1)
        header_row = 6
        for col_idx in range(1, len(headers1) + 1):
            cell = ws1.cell(row=header_row, column=col_idx)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = center_align
            cell.border = thin_border

        # 填入群员统计明细
        for stat in summary.member_stats:
            row_vals = [
                stat.rank,
                stat.nickname,
                stat.message_count,
                stat.ratio,
                stat.total_words,
                stat.first_msg_time or "-",
                stat.last_msg_time or "-"
            ]
            ws1.append(row_vals)
            curr_row = ws1.max_row
            # 格式化
            ws1.cell(row=curr_row, column=1).alignment = center_align
            ws1.cell(row=curr_row, column=2).alignment = left_align
            ws1.cell(row=curr_row, column=3).alignment = right_align
            ws1.cell(row=curr_row, column=3).number_format = "#,##0"
            ws1.cell(row=curr_row, column=4).alignment = right_align
            ws1.cell(row=curr_row, column=4).number_format = "0.0%"
            ws1.cell(row=curr_row, column=5).alignment = right_align
            ws1.cell(row=curr_row, column=5).number_format = "#,##0"
            ws1.cell(row=curr_row, column=6).alignment = center_align
            ws1.cell(row=curr_row, column=7).alignment = center_align
            for col_idx in range(1, len(headers1) + 1):
                c = ws1.cell(row=curr_row, column=col_idx)
                c.font = regular_font
                c.border = thin_border

        # -----------------------------------------------------------------
        # Sheet 2: 24小时时段分布
        # -----------------------------------------------------------------
        ws2 = wb.create_sheet(title="24小时时段分布")
        ws2.views.sheetView[0].showGridLines = True
        headers2 = ["时段", "消息条数", "全天占比"]
        ws2.append(headers2)
        for col_idx in range(1, len(headers2) + 1):
            cell = ws2.cell(row=1, column=col_idx)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = center_align
            cell.border = thin_border

        for h in range(24):
            cnt = summary.hourly_distribution.get(h, 0)
            ratio = (cnt / summary.total_messages) if summary.total_messages > 0 else 0.0
            row_data = [f"{h:02d}:00 - {h+1:02d}:00", cnt, ratio]
            ws2.append(row_data)
            curr_row = ws2.max_row
            ws2.cell(row=curr_row, column=1).alignment = center_align
            ws2.cell(row=curr_row, column=2).alignment = right_align
            ws2.cell(row=curr_row, column=2).number_format = "#,##0"
            ws2.cell(row=curr_row, column=3).alignment = right_align
            ws2.cell(row=curr_row, column=3).number_format = "0.0%"
            for col_idx in range(1, len(headers2) + 1):
                c = ws2.cell(row=curr_row, column=col_idx)
                c.font = regular_font
                c.border = thin_border

        # -----------------------------------------------------------------
        # Sheet 3: 全天消息明细（若提供）
        # -----------------------------------------------------------------
        if messages:
            ws3 = wb.create_sheet(title="全天消息明细")
            ws3.views.sheetView[0].showGridLines = True
            headers3 = ["序号", "发送时间", "群员昵称", "消息类型", "消息内容", "字数"]
            ws3.append(headers3)
            for col_idx in range(1, len(headers3) + 1):
                cell = ws3.cell(row=1, column=col_idx)
                cell.font = header_font
                cell.fill = header_fill
                cell.alignment = center_align
                cell.border = thin_border

            for idx, msg in enumerate(messages, 1):
                ws3.append([
                    idx,
                    msg.time_str,
                    msg.sender_nickname,
                    msg.msg_type,
                    msg.content,
                    msg.word_count
                ])
                curr_row = ws3.max_row
                ws3.cell(row=curr_row, column=1).alignment = center_align
                ws3.cell(row=curr_row, column=2).alignment = center_align
                ws3.cell(row=curr_row, column=3).alignment = left_align
                ws3.cell(row=curr_row, column=4).alignment = center_align
                ws3.cell(row=curr_row, column=5).alignment = left_align
                ws3.cell(row=curr_row, column=6).alignment = right_align
                for col_idx in range(1, len(headers3) + 1):
                    c = ws3.cell(row=curr_row, column=col_idx)
                    c.font = regular_font
                    c.border = thin_border

        # 自动调整各列宽度
        for ws in wb.worksheets:
            for col in ws.columns:
                max_len = 0
                col_letter = get_column_letter(col[0].column)
                for cell in col:
                    val_str = str(cell.value or "")
                    # 中文字符长度计算加权
                    w = sum(2 if ord(ch) > 127 else 1 for ch in val_str)
                    if w > max_len:
                        max_len = w
                ws.column_dimensions[col_letter].width = max(max_len + 3, 12)

        wb.save(output_path)
        return output_path
